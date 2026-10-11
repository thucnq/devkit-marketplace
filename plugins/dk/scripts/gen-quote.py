#!/usr/bin/env python3
"""Generate the client quote workbook (layout of 'QLDN - Báo giá giai đoạn 1') from docs/estimate/items.json.

Sheets: 'Báo giá' (summary, all totals are live formulas), 'Hạng mục' (work breakdown),
'Câu hỏi cần chốt' (open questions). Colors: yellow = cần chốt, orange = tùy chọn, green = tổng.
Usage: gen-quote.py items.json out.xlsx
"""
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("gen-quote: cần openpyxl → pip install openpyxl")

YELLOW = PatternFill("solid", fgColor="FFF2CC")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
GREEN = PatternFill("solid", fgColor="E2EFDA")
HEAD = PatternFill("solid", fgColor="1F3864")
F = "Arial"
VND = '#,##0;-#,##0;"-"'
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(**kw):
    return Font(name=F, size=kw.pop("size", 10), **kw)


def main(src, out):
    d = json.load(open(src, encoding="utf-8"))
    p = {"hourly_rate": 175000, "hours_per_day": 8, "overhead_pct": 0.30, "engineers": 3,
         "working_days_per_month": 22, "usd_vnd": 26000, "ops_fee_month": 0, "ops_note": "",
         "deposit_pct": 0.5, "warranty_months": 2, "uat_weeks": 2, **d.get("params", {})}
    if p["overhead_pct"] > 1:  # accept 30 as 30%
        p["overhead_pct"] /= 100
    infra = d.get("infra") or [{"name": "Cloud", "one_time": 0, "one_time_billed": 0, "monthly": 0}]
    items = d.get("items", [])

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    s = wb.active
    s.title = "Báo giá"
    hm = wb.create_sheet("Hạng mục")
    qs = wb.create_sheet("Câu hỏi cần chốt")

    # ---------------- Hạng mục ----------------
    hdr = ["STT", "Phân hệ", "Hạng mục", "Mô tả / AC", "Điểm (ngày công)", "Loại",
           "Ngày công (gồm QL, BA, QA)", "Thành tiền (VNĐ)", "Ghi chú / cần chốt"]
    for c, h in enumerate(hdr, 1):
        cell = hm.cell(row=1, column=c, value=h)
        cell.font = font(bold=True, color="FFFFFF"); cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    first = 2
    for i, it in enumerate(items):
        r = first + i
        vals = [i + 1, it.get("module", ""), it["name"], it.get("desc", ""), float(it["points"]),
                "Tùy chọn" if it.get("optional") else "Bắt buộc",
                f"=E{r}*(1+'Báo giá'!$C$8)", f"=G{r}*'Báo giá'!$C$7", it.get("question", "")]
        for c, v in enumerate(vals, 1):
            cell = hm.cell(row=r, column=c, value=v)
            cell.font = font(color="0000FF" if c == 5 else "000000")
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = YELLOW if it.get("open") else ORANGE if it.get("optional") else None
        if fill:
            for c in range(1, 10):
                hm.cell(row=r, column=c).fill = fill
        hm.cell(row=r, column=7).number_format = "0.0"
        hm.cell(row=r, column=8).number_format = VND
    last = first + max(len(items), 1) - 1
    tr = last + 1
    hm.cell(row=tr, column=3, value="TỔNG").font = font(bold=True)
    for c, f in [(5, f"=SUM(E{first}:E{last})"), (7, f"=SUM(G{first}:G{last})"), (8, f"=SUM(H{first}:H{last})")]:
        cell = hm.cell(row=tr, column=c, value=f)
        cell.font = font(bold=True); cell.fill = GREEN
        cell.number_format = VND if c == 8 else "0.0"
    for col, w in zip("ABCDEFGHI", [5, 16, 32, 44, 12, 11, 14, 16, 36]):
        hm.column_dimensions[col].width = w
    hm.freeze_panes = "A2"
    E = f"'Hạng mục'!$E${first}:$E${last}"
    T = f"'Hạng mục'!$F${first}:$F${last}"

    # ---------------- Báo giá ----------------
    s["A1"] = d.get("title", "BÁO GIÁ")
    s["A1"].font = font(bold=True, size=14)
    s["A2"] = d.get("subtitle", "")
    s["A2"].font = font(italic=True)
    rows = [
        (4, "THÔNG SỐ ĐẦU VÀO", None, "Ghi chú"),
        (5, "Đơn giá / GIỜ (VNĐ)", p["hourly_rate"], ""),
        (6, "Giờ / ngày công", p["hours_per_day"], "1 điểm ≈ 1 ngày công"),
        (7, "Đơn giá / ngày (VNĐ)", "=C5*C6", "Tự động: giờ/ngày × đơn giá/giờ"),
        (8, "Giám đốc dự án, Phân tích Nghiệp vụ, Tự động Đảm bảo Chất lượng (%)", p["overhead_pct"], ""),
        (9, "Số kỹ sư phần mềm song song", p["engineers"], ""),
        (10, "Tỷ giá USD => VNĐ", p["usd_vnd"], ""),
        (11, "Phí vận hành/tháng", p["ops_fee_month"], p["ops_note"]),
        (12, "Ngày làm việc / tháng", p["working_days_per_month"], "Dùng để quy đổi thời gian"),
        (14, "A. CHI PHÍ PHÁT TRIỂN (build)", None, None),
        (15, "Điểm bắt buộc (ngày công)", f'=SUMIFS({E},{T},"Bắt buộc")', ""),
        (16, "Điểm tùy chọn (ngày công)", f'=SUMIFS({E},{T},"Tùy chọn")', ""),
        (17, "Ngày công (bao gồm cả Giám đốc dự án, Phân tích Nghiệp vụ, Tự động Đảm bảo Chất lượng)", "=C15*(1+C8)", ""),
        (18, "Chi phí build BẮT BUỘC", "=C17*C7", ""),
        (19, "Chi phí build TÙY CHỌN", "=C16*(1+C8)*C7", ""),
        (20, "TỔNG build", "=C18+C19", ""),
        (21, "Thời gian ước tính (tháng)", "=IF(C9*C12=0,0,C17/C9/C12)", "Chỉ tính phần bắt buộc"),
    ]
    inputs = {5, 6, 8, 9, 10, 11, 12}
    for r, label, val, note in rows:
        s.cell(row=r, column=2, value=label).font = font(bold=r in (4, 14, 20))
        if val is not None:
            c = s.cell(row=r, column=3, value=val)
            c.font = font(color="0000FF" if r in inputs else "000000", bold=r == 20)
            c.number_format = "0%" if r == 8 else "0.0" if r in (15, 16, 17, 21) else VND
            c.border = BOX
            if r in (20, 21):
                c.fill = GREEN
        if note:
            s.cell(row=r, column=4, value=note).font = font(italic=True, color="595959")

    n = len(infra)
    cols = [get_column_letter(3 + i) for i in range(n)]
    s["B23"] = "B. HẠ TẦNG"; s["B23"].font = font(bold=True)
    s["B30"] = "C. Quy trình bàn giao, bảo trì và thanh toán"; s["B30"].font = font(bold=True)
    for i, inf in enumerate(infra):
        col = cols[i]
        for hr in (24, 31):
            h = s[f"{col}{hr}"]; h.value = inf["name"]; h.font = font(bold=True); h.fill = GREEN
        for r, v in [(25, inf.get("one_time", 0)), (26, inf.get("one_time_billed", inf.get("one_time", 0))),
                     (27, inf.get("monthly", 0)), (28, f"={col}25+{col}27*36"), (29, f"={col}25+{col}27*60"),
                     (32, f"=$C$18*{p['deposit_pct']}+{col}26"), (33, 0), (34, f"=$C$18*(1-{p['deposit_pct']})"),
                     (35, 0), (36, f"={col}32+{col}34")]:
            c = s[f"{col}{r}"]; c.value = v; c.number_format = VND; c.border = BOX
            c.font = font(color="0000FF" if r in (25, 26, 27) else "000000", bold=r == 36)
            if r == 36:
                c.fill = GREEN
    labels = {25: "Một lần (phần cứng + tài khoản + setup)", 26: "Trong đó thu qua báo giá (một lần)",
              27: "Định kỳ / tháng", 28: "TCO 3 năm", 29: "TCO 5 năm",
              32: "Ngày ký hợp đồng",
              33: f"Sau khi phát triển xong một số nhóm tính năng, sản phẩm được triển khai để khách hàng nghiệm thu trong {p['uat_weeks']} tuần. Lỗi hoặc chỉnh sửa nhỏ không đổi nghiệp vụ được sửa và triển khai trong lần kế tiếp",
              34: "Ngày kết thúc giai đoạn",
              35: f"Trong {p['warranty_months']} tháng bảo hành, lỗi do đội phát triển được xử lý sớm nhất có thể và giải thích nguyên nhân trong 48 giờ",
              36: "Tổng chi phí"}
    for r, t in labels.items():
        c = s.cell(row=r, column=2, value=t)
        c.font = font(bold=r == 36); c.alignment = Alignment(wrap_text=True, vertical="top")

    legend = [(38, "CHÚ THÍCH", None), (39, "Ô vàng = cần chốt (xem sheet Câu hỏi cần chốt)", YELLOW),
              (40, "Cam = tùy chọn", ORANGE), (41, "Xanh = tổng/kết quả", GREEN),
              (42, "Chữ xanh dương = thông số nhập tay, có thể sửa; chữ đen = công thức", None)]
    for r, t, fill in legend:
        c = s.cell(row=r, column=2, value=t); c.font = font(bold=r == 38)
        if fill:
            c.fill = fill
    s.column_dimensions["A"].width = 3
    s.column_dimensions["B"].width = 58
    for col in cols:
        s.column_dimensions[col].width = 20
    s.column_dimensions[get_column_letter(3 + n)].width = 44 if n < 2 else 20
    if n < 2:
        s.column_dimensions["D"].width = 44

    # ---------------- Câu hỏi cần chốt ----------------
    qh = ["STT", "Câu hỏi", "Hạng mục liên quan", "Vì sao cần chốt", "Mặc định nếu không trả lời", "Trả lời của khách"]
    for c, h in enumerate(qh, 1):
        cell = qs.cell(row=1, column=c, value=h)
        cell.font = font(bold=True, color="FFFFFF"); cell.fill = HEAD
    qrows = [(it.get("question", ""), it["name"], it.get("why", "Ảnh hưởng điểm ước lượng"), it.get("default", ""))
             for it in items if it.get("open")]
    qrows += [(q["q"], q.get("item", ""), q.get("why", ""), q.get("default", "")) for q in d.get("questions", [])]
    for i, (q, item, why, default) in enumerate(qrows):
        for c, v in enumerate([i + 1, q, item, why, default, ""], 1):
            cell = qs.cell(row=i + 2, column=c, value=v)
            cell.font = font(); cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c == 6:
                cell.fill = YELLOW
    for col, w in zip("ABCDEF", [5, 48, 24, 32, 32, 32]):
        qs.column_dimensions[col].width = w

    wb.save(out)
    # Same arithmetic as the sheet formulas, so the chat summary never relies on LLM math.
    mand = sum(float(i["points"]) for i in items if not i.get("optional"))
    opt = sum(float(i["points"]) for i in items if i.get("optional"))
    day_rate = p["hourly_rate"] * p["hours_per_day"]
    md = mand * (1 + p["overhead_pct"])
    months = md / p["engineers"] / p["working_days_per_month"] if p["engineers"] else 0
    print(f"gen-quote: {len(items)} hạng mục, {len(qrows)} câu hỏi → {out}")
    print(f"  điểm bắt buộc {mand:g} · tùy chọn {opt:g} · ngày công (gồm QL/BA/QA) {md:g}")
    print(f"  build bắt buộc {md * day_rate:,.0f} VNĐ · tùy chọn {opt * (1 + p['overhead_pct']) * day_rate:,.0f} VNĐ · ~{months:.1f} tháng với {p['engineers']} kỹ sư")
    for inf in infra:
        print(f"  {inf['name']}: một lần {inf.get('one_time', 0):,.0f} · tháng {inf.get('monthly', 0):,.0f} · TCO 3 năm {inf.get('one_time', 0) + inf.get('monthly', 0) * 36:,.0f}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
